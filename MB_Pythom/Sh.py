from tkinter import *
from tkinter import ttk
from tkinter import scrolledtext  

import sys

import os
import time
# import shutil
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import pyModbusTCP
from pyModbusTCP.client import ModbusClient

# global scanIP, scanId, tabFile, rowSt, rowCnt, colSt, maxCmdInRow
# import wxPy
# from wxPy import cmdUpdate

def sendCmdBlock(tabFile, rowSt, rowCnt, colSt, maxCmdInRow):
    #Имя файла с командами для настройки ПГС
    file_path = 'PgsSetup.xlsx'
    # позиция и размер таблицы
    start_row = rowSt
    row_count = 3
    start_col = 3
    col_count = 12

    # if __name__ == '__main__':
    cmdText_l = "Команды настройки:\n"

    cells_in_cmd = 3
    scan_IP = '192.168.70.33'
    scan_ID = 5
    if len(sys.argv) > 5:
        file_path = sys.argv[1]
        start_row = int(sys.argv[2])
        row_count = int(sys.argv[3])
        start_col = int(sys.argv[4])
        col_count = int(sys.argv[5]) * cells_in_cmd
        if len(sys.argv) > 6:
            scan_IP = sys.argv[6]
        if len(sys.argv) > 7:
            scan_ID = int(sys.argv[7])

        # print(f'start_row = {start_row}, row_count = {row_count}')
        # print(f'start_col = {start_col}, col_count = {col_count}')
    else:
        file_path = tabFile
        start_row = int(rowSt)
        row_count = int(rowCnt)
        start_col = int(colSt)
        col_count = int(maxCmdInRow) * cells_in_cmd
    
    last_row = start_row + row_count - 1
    err = True
    # Проверяем, существует ли файл
    if os.path.exists(file_path):
        # print(f"Есть Файл {file_path}")
        workbook = load_workbook(file_path)
        page = workbook.active
        
        client = ModbusClient(host = scan_IP, port = 502, unit_id = scan_ID, timeout = 0.5, auto_open = True)
    #    ModbusClient(host="192.168.56.101", port=502, debug= False, unit_id=5, auto_open=True, auto_close=True)
        curr_row = start_row
        for row in page.iter_rows(min_row = start_row, max_row = last_row, values_only=True):
            for i in range(start_col, start_col + col_count, cells_in_cmd):
                curr_value = page.cell(row = curr_row, column = i + 1).value
                # print(curr_row, i, curr_value)
                if (curr_value == None):
                    break
                else:
                    aVal = [row[i + 0], row[i + 1], row[i + 2]]
                    
                    aCmd = [row[i + 0], row[i + 1], row[i + 2]]
                    for i in range(0, 2):
                        if (aVal[i] < 0):
                            aCmd[i] += 65536
                    err = client.write_multiple_registers(2901, aCmd)
                    if (err == False):
                        # print(f"{aVal} Нет ответа!")
                        cmd = f"{aVal} Нет ответа!"
                    else:
                        # print(aVal)
                        cmd = f"{aVal}"
                    # print(cmd)
                    cmdText_l += (cmd + "\n")
                    if (err == False):
                        break
            curr_row = curr_row + 1
            if (err == False):
                break
        client.close()
    else:
        print(f"Нет Файла {file_path}")
    return cmdText_l


# === Добавлено: карты соответствий для экспорта в Excel ===
CHANNEL_TYPE_MAP = {
    0x30: "КТВ",
    0x31: "D_in",
    0x32: "R_Out",
    0x33: "интерком",
    0x34: "A_in",
    0x35: "TEU",
}

MODE_MAP = {
    0x0: "НО",
    0x1: "НЗ",
    0x2: "РО",
    0x3: "РЗ",
}

# === Добавлено: вспомогательные функции форматирования для экспорта ===
def _format_hex(value):
    return f"0x{value:02X}"


def _map_value(value, mapping):
    return mapping.get(value, "Неизвестно")
# === Конец добавления: вспомогательные функции форматирования для экспорта ===


# === Добавлено: экспорт постов в Excel ===
def export_posts_to_excel(post_count, filename, scan_ip, scan_id, output_callback=None):
    if not filename.lower().endswith(".xlsx"):
        filename = f"{filename}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Экспорт"
    header_row = ["Пост", "Тип", "Адрес", "Режим"]
    sheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="D9D9D9")
    header_font = Font(bold=True)
    border_side = Side(style="thin", color="000000")
    cell_border = Border(
        left=border_side,
        right=border_side,
        top=border_side,
        bottom=border_side,
    )

    posts_per_row = 5
    columns_per_post = 4
    spacer_columns = 1
    max_columns = (posts_per_row * columns_per_post) + ((posts_per_row - 1) * spacer_columns)

    column_widths = {
        1: 10,
        2: 10,
        3: 10,
        4: 12,
    }
    for post_index in range(posts_per_row):
        base_column = (post_index * (columns_per_post + spacer_columns)) + 1
        for offset, width in column_widths.items():
            sheet.column_dimensions[chr(64 + base_column + offset - 1)].width = width
        spacer_column = base_column + columns_per_post
        if post_index < posts_per_row - 1:
            sheet.column_dimensions[chr(64 + spacer_column)].width = 2

    post_fills = [
        PatternFill("solid", fgColor="E8F1FF"),
        PatternFill("solid", fgColor="E9F7EF"),
    ]

    def _write_cell(row, column, value, fill=None, font=None):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border
        if fill:
            cell.fill = fill
        if font:
            cell.font = font
        return cell

    def _write_spacer_cell(row, column):
        return _write_cell(row, column, None, fill=PatternFill("solid", fgColor="FFFFFF"))

    client = ModbusClient(
        host=scan_ip,
        port=502,
        unit_id=scan_id,
        timeout=0.5,
        auto_open=True,
    )

    def _get_post_rows(post_index):
        register_start = 10500 + post_index
        registers = client.read_holding_registers(register_start, 24)
        if registers is None:
            if output_callback:
                output_callback(f"Пост {post_index}: нет ответа")
            return [["", "-", "-", "Нет ответа"]] + [["", "-", "-", "-"] for _ in range(11)]

        raw_bytes = []
        for register in registers:
            raw_bytes.append(register & 0xFF)
            raw_bytes.append((register >> 8) & 0xFF)

        rows = []
        for channel_index in range(1, 13):
            offset = (channel_index - 1) * 4
            channel_bytes = raw_bytes[offset:offset + 4]
            if len(channel_bytes) < 4:
                break

            if channel_bytes[0] == 0x00:
                row_data = ["", "-", "-", "-"]
            else:
                type_label = _map_value(channel_bytes[0], CHANNEL_TYPE_MAP)
                mode_label = _map_value(channel_bytes[2], MODE_MAP)
                address_value = channel_bytes[1]

                if channel_bytes[0] == 0x30:
                    mode_label = f"№ {channel_bytes[2]}"
                elif channel_bytes[0] == 0x35:
                    if channel_bytes[2] == 0xFF:
                        mode_label = "-"
                        address_value = "-"
                    else:
                        mode_label = channel_bytes[2]
                elif channel_bytes[0] == 0x34:
                    adjusted_value = channel_bytes[2] - 127
                    mode_label = f"0 = {adjusted_value}"

                row_data = ["", type_label, address_value, mode_label]
            rows.append(row_data)

        while len(rows) < 12:
            rows.append(["", "-", "-", "-"])
        return rows

    current_row = 1
    for block_start in range(1, post_count + 1, posts_per_row):
        posts_in_block = min(posts_per_row, post_count - block_start + 1)

        for index in range(posts_in_block):
            col_start = 1 + index * (columns_per_post + spacer_columns)
            for offset, value in enumerate(header_row):
                _write_cell(current_row, col_start + offset, value, fill=header_fill, font=header_font)
            if index < posts_in_block - 1:
                _write_spacer_cell(current_row, col_start + columns_per_post)
        current_row += 1

        for index in range(posts_in_block):
            post_index = block_start + index
            col_start = 1 + index * (columns_per_post + spacer_columns)
            row_fill = post_fills[(post_index - 1) % len(post_fills)]
            _write_cell(current_row, col_start, f"Пост {post_index}", fill=row_fill)
            for offset in range(1, columns_per_post):
                _write_cell(current_row, col_start + offset, "", fill=row_fill)
            if index < posts_in_block - 1:
                _write_spacer_cell(current_row, col_start + columns_per_post)
        current_row += 1

        post_rows = [_get_post_rows(block_start + index) for index in range(posts_in_block)]
        for row_index in range(12):
            for index in range(posts_in_block):
                post_index = block_start + index
                col_start = 1 + index * (columns_per_post + spacer_columns)
                row_fill = post_fills[(post_index - 1) % len(post_fills)]
                row_data = post_rows[index][row_index]
                for offset, value in enumerate(row_data):
                    _write_cell(current_row, col_start + offset, value, fill=row_fill)
                if index < posts_in_block - 1:
                    _write_spacer_cell(current_row, col_start + columns_per_post)
            current_row += 1

        for col_index in range(1, (posts_in_block * (columns_per_post + spacer_columns))):
            if (col_index % (columns_per_post + spacer_columns)) == 0:
                _write_spacer_cell(current_row, col_index)
            else:
                _write_cell(current_row, col_index, "", fill=PatternFill("solid", fgColor="FFFFFF"))
        current_row += 1

    client.close()
    workbook.save(filename)
    return filename
# === Конец добавления: экспорт постов в Excel ===
